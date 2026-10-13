# 12_figures_and_tables.py
# Collects every manuscript figure redrawn from v25 into one folder and rebuilds
# Table S1 (country estimates) and Table S2 (region assignment) from the v25 outputs.
# Count-based ("crude") prevalence columns are not used.
#
# Usage (from D:\early-weaning-analysis), after 02, 03, 04 and 10_country_map have been run on v25:
#   python analysis\12_figures_and_tables.py
#   python analysis\12_figures_and_tables.py --out .\figures_v25

import argparse
import shutil
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

ap = argparse.ArgumentParser()
ap.add_argument("--extract", default=r".\02_extract_v25")
ap.add_argument("--survival", default=r".\03_survival_v25")
ap.add_argument("--comparative", default=r".\04_comparative_v25")
ap.add_argument("--map", default=r".\r30_country_map_v25")
ap.add_argument("--out", default=r".\figures_v25")
a = ap.parse_args()

S3, S4, MAP = Path(a.survival), Path(a.comparative), Path(a.map)
out = Path(a.out)
out.mkdir(parents=True, exist_ok=True)

# v24 submission name -> v25 source
FIGURES = [
    ("Figure2_hazard_ratios.png", S3 / "cox_models" / "hazard_ratios_forest_plot.png"),
    ("Figure3_cessation_before_6_months.png", S3 / "tables" / "early_weaning_top20_countries.png"),
    ("Figure4_curves_by_country.png", S3 / "survival_curves" / "km_by_country.png"),
    ("Figure5_curves_by_region.png", S3 / "survival_curves" / "km_by_region.png"),
    ("Figure6_curves_by_residence.png", S3 / "survival_curves" / "km_by_residence.png"),
    ("Figure7_regional_patterns.png", S4 / "regional" / "regional_patterns.png"),
    ("Figure8_country_rankings.png", S4 / "rankings" / "country_rankings_visualization.png"),
    ("Figure9_never_breastfed_country.png", S4 / "inequalities" / "never_breastfed_by_country.png"),
    ("Figure10_never_breastfed_region.png", S4 / "inequalities" / "never_breastfed_by_region.png"),
    ("Figure11_urban_rural_gap.png", S4 / "inequalities" / "urban_rural_inequality.png"),
    ("Map_country_sample_n_children.png", MAP / "r30_country_map_n_children.png"),
    ("Map_country_sample_n_children.tif", MAP / "r30_country_map_n_children.tif"),
    ("Map_country_sample_data.xlsx", MAP / "r30_country_map_data.xlsx"),
]
index = []
for name, src in FIGURES:
    if src.exists():
        shutil.copy2(src, out / name)
        index.append(f"{name}\t<- {src}")
    else:
        index.append(f"{name}\tMISSING: {src}")
        print(f"MISSING: {src}")

# -------------------- Tables S1 and S2 --------------------
combined = pd.read_csv(Path(a.extract) / "combined_breastfeeding_data.csv",
                       usecols=["country", "source_zip"], low_memory=False)
n_children, n_files = len(combined), combined["source_zip"].nunique()
n_surveys_by_country = combined.groupby("country")["source_zip"].nunique()
rank = pd.read_excel(S4 / "rankings" / "country_rankings_complete.xlsx")
rank["n_surveys"] = rank["country_code"].map(n_surveys_by_country)


def write_table(path, title, note, headers, data_rows, footer_rows=()):
    wb = Workbook()
    ws = wb.active
    ws.title = path.stem.split("_")[1] if "_" in path.stem else "Table"
    ws["A1"] = title
    ws["A1"].font = Font(bold=True)
    ws["A2"] = note
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[2].height = 90
    for j, h in enumerate(headers, 1):
        c = ws.cell(row=4, column=j, value=h)
        c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True, vertical="bottom")
    for i, r in enumerate(data_rows, 5):
        for j, v in enumerate(r, 1):
            ws.cell(row=i, column=j, value=v)
    i = 5 + len(data_rows)
    for r in footer_rows:
        for j, v in enumerate(r, 1):
            ws.cell(row=i, column=j, value=v)
        i += 1
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["C"].width = 26
    for col in "DEFGHIJK":
        ws.column_dimensions[col].width = 16
    ws.freeze_panes = "A5"
    wb.save(path)


r1 = lambda x: None if pd.isna(x) else round(float(x), 1)
r0 = lambda x: None if pd.isna(x) else int(round(float(x)))

s1 = rank.sort_values(["median_duration_months", "country_name"], ascending=[False, True])
write_table(
    out / "Table_S1_country_estimates.xlsx",
    "Table S1. Breastfeeding duration and cessation by country",
    f"Survey-weighted estimates (v005/1,000,000) from {n_files} DHS country-survey datasets in "
    f"{len(rank)} countries, {n_children:,} children. Medians and percentages come from the survival "
    "curve: the Turnbull estimator where children coded m4 = 93 (breastfeeding ended, duration not "
    "recorded) are present, and the Kaplan-Meier estimator otherwise. A median of 36 months means the "
    "curve did not reach 50% within follow-up and is reported at that limit. Never breastfed is DHS "
    "code m4 = 94. Country is taken from the DHS country code in each file (v000); the 1998-99 "
    "Indian state (NFHS-2) files, which duplicate the national survey, and the Dominican Republic "
    "batey special surveys are excluded.",
    ["DHS code", "Country", "Region", "Surveys (n)", "Children (n)", "Left-censored (n)",
     "Median duration (months)", "Never breastfed (%)", "Ceased before 6 months (%)",
     "Still breastfeeding at 12 months (%)", "Still breastfeeding at 24 months (%)"],
    [[r.country_code, r.country_name, r.region, r0(r.n_surveys), r0(r.n_children),
      r0(r.n_left_censored), r0(r.median_duration_months), r1(r.pct_never_breastfed),
      r1(r.pct_weaned_before_6mo), r1(r.pct_still_bf_at_12mo), r1(r.pct_still_bf_at_24mo)]
     for r in s1.itertuples()])

s2 = rank.sort_values(["region", "country_code"])
write_table(
    out / "Table_S2_region_assignment.xlsx",
    "Table S2. Assignment of DHS countries to the ten analytic regions",
    "Regions are the pragmatic analytic groupings used in this study, based on geographic "
    "proximity; they are not a standard international classification such as UN M49 or the World "
    "Bank regions, and regional estimates should be read accordingly. The assignment is fixed in the "
    "analysis code. Codes are DHS country codes, not ISO codes (e.g. MD = Madagascar, MB = Moldova, "
    f"NI = Niger, BU = Burundi). Estimates are survey-weighted and cover all children (n = {n_children:,}).",
    ["DHS code", "Country", "Region", "Surveys (n)", "Children (n)", "Left-censored (n)",
     "Median duration (months)", "Still breastfeeding at 24 months (%)"],
    [[r.country_code, r.country_name, r.region, r0(r.n_surveys), r0(r.n_children),
      r0(r.n_left_censored), r0(r.median_duration_months), r1(r.pct_still_bf_at_24mo)]
     for r in s2.itertuples()],
    footer_rows=[[f"Total ({len(rank)} countries, {rank['region'].nunique()} regions)", None, None,
                  int(rank["n_surveys"].sum()), int(rank["n_children"].sum())]])
index += ["Table_S1_country_estimates.xlsx\t<- rebuilt from 04 country_rankings_complete.xlsx",
          "Table_S2_region_assignment.xlsx\t<- rebuilt from 04 country_rankings_complete.xlsx"]

(out / "INDEX.txt").write_text(
    "Figures and supplementary tables redrawn from v25. File names follow the v24 submission\n"
    "numbering; renumber to the current manuscript as needed. Figure 1 (methods workflow) is a\n"
    "drawn diagram, not regenerated: check any counts it shows against r31_key_numbers_v25.xlsx.\n\n"
    + "\n".join(index) + "\n", encoding="utf-8")
print("\n".join(index))
