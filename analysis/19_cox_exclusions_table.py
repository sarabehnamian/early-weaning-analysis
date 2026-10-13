# 19_cox_exclusions_table.py
# Table S8: every country-survey in the v25 extraction that contributes no children to the
# primary Cox model (cox_common.cox_sample), with the reason and the number of children.
# A survey (country x survey year, the Cox stratum) is excluded when it has
#   - no retrospective durations: no child with a recorded cessation after month 0;
#   - no wealth index: wealth_quintile missing for every child;
#   - no PSU: v021 missing for every child;
# or (checked for completeness) no positive weight, no residence or no maternal education.
# All reasons that apply are listed. Surveys that keep some children are not listed.
#
# Usage (from D:\early-weaning-analysis):  python analysis\19_cox_exclusions_table.py
# Output: v25_extras\Table_S8_cox_exclusions.xlsx

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from cox_common import load

OUT = Path(r".\v25_extras")
OUT.mkdir(exist_ok=True)
RANK = Path(r".\04_comparative_v25\rankings\country_rankings_complete.xlsx")

df = load()
key = ["country", "survey_year"]
df["pos_w"] = df["weight"] > 0
df["retro_event"] = df["pos_w"] & (df["event"] == 1) & (df["duration_months"] > 0)
for c in ["wealth_quintile", "v021", "urban", "mother_educ"]:
    df[f"has_{c}"] = df[c].notna()
df["in_cox"] = (df["pos_w"] & (df["censor_type"] != "left") & df["has_wealth_quintile"]
                & df["has_v021"] & df["has_urban"] & df["has_mother_educ"])

g = df.groupby(key)
s = g[["pos_w", "retro_event", "has_wealth_quintile", "has_v021", "has_urban",
       "has_mother_educ", "in_cox"]].sum()
s["n_children"] = g.size()
s["files"] = g["source_zip"].agg(lambda x: ", ".join(sorted(set(map(str, x)))))
# retrospective surveys only enter the Cox sample at all
s["in_cox"] = s["in_cox"].where(s["retro_event"] > 0, 0)
ex = s[s["in_cox"] == 0].reset_index()

REASONS = [("pos_w", "no positive survey weight"),
           ("retro_event", "no retrospective durations"),
           ("has_wealth_quintile", "no wealth index"),
           ("has_v021", "no PSU (v021)"),
           ("has_urban", "no residence"),
           ("has_mother_educ", "no maternal education")]
ex["reason"] = ex.apply(lambda r: "; ".join(t for c, t in REASONS if r[c] == 0), axis=1)

names = pd.read_excel(RANK, usecols=["country_code", "country_name"]).drop_duplicates()
ex = ex.merge(names, left_on="country", right_on="country_code", how="left")
cox_countries = set(s[s["in_cox"] > 0].index.get_level_values("country"))
ex["other_surveys_in_cox"] = ex["country"].map(lambda c: "Yes" if c in cox_countries else "No")
ex = ex.sort_values(["country_name", "survey_year"])

n_all, n_cox = len(df), int(s["in_cox"].sum())
n_surv, n_ex, n_ex_children = len(s), len(ex), int(ex["n_children"].sum())
whole = sorted(ex.loc[ex["other_surveys_in_cox"] == "No", "country_name"].dropna().unique())

wb = Workbook()
ws = wb.active
ws.title = "S8"
headers = ["DHS code", "Country", "Survey year", "DHS file(s)", "Children (n)",
           "Reason for exclusion", "Country has other surveys in Cox model"]
ws["A1"] = "Table S8. Country-surveys excluded from the Cox model"
ws["A1"].font = Font(bold=True)
ws["A2"] = (
    f"Of {n_surv} country-surveys (country x survey year, the Cox strata) in the analytic dataset, "
    f"{n_ex} contribute no children to the stratified Cox model ({n_ex_children:,} of {n_all:,} "
    f"children). A survey is excluded when it has no retrospective durations (no child with a "
    f"recorded cessation after month 0, so the survey cannot inform the within-stratum hazard), "
    f"no household wealth index, or no primary sampling unit (v021) for clustered standard errors; "
    f"all reasons that apply are listed. Children are counted before any other exclusion. "
    f"Within included surveys, children with m4 = 93 (left-censored) or missing covariates are "
    f"also dropped, giving a Cox sample of {n_cox:,}. "
    + (f"Countries with no survey in the Cox model: {', '.join(whole)}." if whole else ""))
ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
ws.row_dimensions[2].height = 120
for j, h in enumerate(headers, 1):
    c = ws.cell(row=4, column=j, value=h)
    c.font = Font(bold=True)
    c.alignment = Alignment(wrap_text=True, vertical="bottom")
i = 5
for r in ex.itertuples():
    for j, v in enumerate([r.country, r.country_name, int(r.survey_year), r.files,
                           int(r.n_children), r.reason, r.other_surveys_in_cox], 1):
        ws.cell(row=i, column=j, value=v)
    i += 1
ws.cell(row=i, column=1, value=f"Total ({n_ex} surveys)").font = Font(bold=True)
ws.cell(row=i, column=5, value=n_ex_children).font = Font(bold=True)
for col, w in zip("ABCDEFG", [10, 24, 12, 28, 14, 44, 16]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "A5"
wb.save(OUT / "Table_S8_cox_exclusions.xlsx")

print(f"{n_ex} of {n_surv} surveys excluded, {n_ex_children:,} children; Cox sample {n_cox:,}")
print(ex["reason"].value_counts().to_string())
print("Countries with no survey in Cox:", ", ".join(whole) or "none")
